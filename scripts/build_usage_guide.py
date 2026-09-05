from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "个性化社交媒体写作Agent_详细使用手册.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "596775"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
BLACK = "1A1A1A"
BORDER = "B8C2CC"


def set_run_font(run, size=None, bold=None, color=BLACK, italic=None, mono=False):
    latin = "Consolas" if mono else "Calibri"
    east_asia = "Microsoft YaHei"
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    assert sum(widths) == 9360
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def style_table(table, header=True, first_col_bold=False):
    table.style = "Table Grid"
    if header and table.rows:
        set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if header and r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_run_font(run, size=9.3, bold=(header and r_idx == 0) or (first_col_bold and c_idx == 0))


def add_table(doc, headers, rows, widths, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
    set_table_geometry(table, widths)
    style_table(table, first_col_bold=first_col_bold)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_numbering_definition(doc, abstract_id, num_id, fmt, text):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    level.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    level.append(jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def add_list_item(doc, text, ordered=False, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "101" if ordered else "100")
    num_pr.extend([ilvl, num_id])
    ppr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        set_run_font(p.add_run(text))
    return p


def add_callout(doc, label, text, fill=CALLOUT):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.14)
    p.paragraph_format.line_spacing = 1.2
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    r1 = p.add_run(label + "  ")
    set_run_font(r1, bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(text))
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.keep_together = True
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    ppr.append(shd)
    run = p.add_run(text)
    set_run_font(run, size=8.6, color="28323C", mono=True)
    return p


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, cached, fld_end])
    set_run_font(run, size=8.5, color=MUTED)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Title"].font.name = "Calibri"
    styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Title"].font.size = Pt(30)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Subtitle"].font.name = "Calibri"
    styles["Subtitle"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Subtitle"].font.size = Pt(15)
    styles["Subtitle"].font.color.rgb = RGBColor.from_string(DARK_BLUE)


def setup_page(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run("个性化社交媒体写作 Agent｜详细使用手册"), size=8.5, color=MUTED)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(0)
        set_run_font(fp.add_run("内部操作指南  ·  "), size=8.5, color=MUTED)
        add_page_number(fp)


def build():
    doc = Document()
    setup_styles(doc)
    setup_page(doc)
    add_numbering_definition(doc, 100, 100, "bullet", "•")
    add_numbering_definition(doc, 101, 101, "decimal", "%1.")

    # Editorial cover
    for _ in range(4):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(kicker.add_run("OPERATOR GUIDE · V1.1"), size=10, bold=True, color=BLUE)
    title = doc.add_paragraph("个性化社交媒体写作 Agent", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    subtitle = doc.add_paragraph("从创建角色、投喂历史文章到 RAG 检索、生成、评审与反馈闭环", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.left_indent = Inches(0.55)
    lead.paragraph_format.right_indent = Inches(0.55)
    set_run_font(lead.add_run("一份面向非技术使用者与项目运营者的完整操作手册"), size=11.5, color=MUTED, italic=True)
    for _ in range(4):
        doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("适用项目：SocialPostEditor\n运行模式：mock / DeepSeek / OpenAI\n更新日期：2026 年 9 月 4 日"), size=10, color=MUTED)
    doc.add_page_break()

    add_heading(doc, "阅读地图", 1)
    add_callout(doc, "一句话理解", "先从当前角色过去的文章中找到最相关的参考，再在角色边界和可信事实范围内生成内容；生成后由独立评审器检查，不合格则定向修改。")
    add_table(doc, ["如果你想……", "直接阅读"], [
        ("第一次跑起来", "第 2、3 节"),
        ("创建新的个人或品牌角色", "第 4 节"),
        ("导入过去的文章", "第 5 节"),
        ("理解画像、RAG 和生成", "第 6–9 节"),
        ("让系统越来越像本人", "第 10 节"),
        ("排查问题或验证系统", "第 12–13 节"),
    ], [2700, 6660], first_col_bold=True)

    add_heading(doc, "1. 系统到底在做什么", 1)
    add_body(doc, "普通模型只看到本次要求，容易写出通用营销腔。这个 Agent 会先组装一份受控资料包，然后再生成。资料包的各部分有不同职责，不能互相替代。")
    add_table(doc, ["组件", "回答的问题", "作用"], [
        ("角色边界", "这个身份是谁，不能说什么？", "阻止串号、虚构身份和越界承诺"),
        ("作者画像", "这个角色通常怎么思考和表达？", "保持长期语气、结构和价值观"),
        ("RAG", "这次应该回忆哪些历史文章？", "按当前主题动态检索参考"),
        ("写作任务", "这次要写什么、给谁看？", "定义平台、受众、目标、格式和长度"),
        ("proof_points", "本次有哪些事实可以使用？", "限制数字、案例、客户结果和具体主张"),
        ("人工反馈", "本人以前如何修改相似初稿？", "把修改偏好沉淀成可复用记忆"),
        ("自动评审", "初稿是否达到发布标准？", "检查事实、照抄、模板感、营销强度与 CTA"),
    ], [1750, 3000, 4610], first_col_bold=True)
    add_callout(doc, "最重要的边界", "历史文章不是自动可信的事实库。历史文章中的数字和案例，除非重新写进本次 proof_points，否则不应被当成当前可公开事实。")

    add_heading(doc, "2. 第一次启动", 1)
    add_heading(doc, "2.1 安装依赖", 2)
    add_body(doc, "在项目目录 C:\\alignment\\social\\SocialPostEditor 打开 PowerShell，然后执行：")
    add_code(doc, "cd C:\\alignment\\social\\SocialPostEditor\n.\\.venv\\Scripts\\Activate.ps1\npip install -r requirements.txt")
    add_heading(doc, "2.2 启动服务", 2)
    add_code(doc, "uvicorn app.main:app --reload")
    add_body(doc, "浏览器打开 http://127.0.0.1:8000/docs。这里是 FastAPI 自动生成的 Swagger 操作页面，可以直接填写 JSON、上传文件和查看结果。")
    add_heading(doc, "2.3 检查运行状态", 2)
    add_body(doc, "在 Swagger 中展开 GET /health，点击 Try it out，再点击 Execute。正常结果应显示 status=ok，并列出当前模型供应商、模型和检索模式。")
    add_code(doc, '{\n  "status": "ok",\n  "llm_mode": "live",\n  "llm_provider": "deepseek",\n  "model": "deepseek-v4-flash",\n  "retrieval_mode": "hybrid"\n}')
    add_callout(doc, "密钥安全", "密钥保存在项目根目录的 .env 中，并已被 .gitignore 排除。不要把 .env 发给他人，不要提交到 Git，也不要在截图或日志中暴露密钥。")

    add_heading(doc, "3. 推荐的完整使用顺序", 1)
    for item in (
        "创建角色并记住返回的 role_id。",
        "批量导入该角色过去真实发布的文章，由系统自动拆分、标注和去重，再人工纠错。",
        "调用画像重建，只让高真实性样本参与作者画像。",
        "先调用检索接口，检查本次主题会取回哪些历史文章。",
        "提交生成任务，明确 proof_points、禁用表达和 CTA。",
        "在 review-inbox 中查看审核摘要，编辑并明确批准或拒绝内容。",
        "保存“初稿—终稿—修改原因”，再人工决定是否准入长期记忆。",
        "积累一批任务后，根据真实效果调整标签、检索权重或模型。",
    ):
        add_list_item(doc, item, ordered=True)
    add_callout(doc, "推荐习惯", "把 /retrieve 作为生成前的固定检查。如果取回的文章明显不相关，应先修正标签和任务描述，而不是让模型带着错误参考继续写。")

    doc.add_page_break()
    add_heading(doc, "4. 创建角色", 1)
    add_body(doc, "角色不是昵称，而是一套独立的数据空间。每个角色拥有自己的文章、画像、反馈和生成记录。品牌官方账号、创始人个人账号、产品账号应尽量拆成不同角色。")
    add_heading(doc, "4.1 在 Swagger 中创建", 2)
    add_body(doc, "找到 POST /roles，点击 Try it out，填写：")
    add_code(doc, '{\n  "name": "Hao｜AI 创业者个人账号",\n  "description": "面向创业者、品牌负责人和 AI 从业者，分享 AI 产品与 GEO 实践。",\n  "identity_rules": "只描述真实经历；不得虚构客户、融资、收入或效果。",\n  "default_generate_params": {\n    "platform": "LinkedIn",\n    "language": "zh-CN",\n    "audience": "品牌与 AI 从业者",\n    "tone": "克制、具体、基于证据",\n    "banned_phrases": ["绝对领先", "颠覆行业"]\n  }\n}')
    add_body(doc, "提交后会返回 id。假设返回 id=3，那么以后导入文章、构建画像、生成和保存反馈都要使用 role_id=3。")
    add_heading(doc, "4.2 角色字段怎么写", 2)
    add_table(doc, ["字段", "填写方法"], [
        ("name", "可识别的账号身份，例如“品牌官方｜中文”或“创始人｜个人观点”"),
        ("description", "受众、核心主题、业务范围、传播任务以及不处理的领域"),
        ("identity_rules", "不可虚构的身份、经历、客户、数字、承诺和法律/合规边界"),
        ("default_generate_params", "角色长期复用的平台、受众、语气、长度和禁用词"),
        ("role_id", "系统返回的角色编号；后续所有资料都通过它隔离"),
    ], [1800, 7560], first_col_bold=True)
    add_heading(doc, "4.3 什么时候应该拆角色", 2)
    for item in (
        "公司品牌账号与创始人个人账号：身份和第一人称经历不同。",
        "不同产品线：受众、事实边界和内容目标明显不同。",
        "中英文账号：如果价值观一致可以先共用角色；如果表达身份或市场定位不同则拆分。",
        "客户代运营账号：每个客户必须独立角色，不能共享文章和反馈。",
    ):
        add_list_item(doc, item)
    add_heading(doc, "4.4 更新与克隆角色", 2)
    add_body(doc, "使用 PATCH /roles/{role_id} 修改身份规则或角色默认生成参数。需要同一作者管理另一个平台时，使用 POST /roles/{role_id}/clone；它复制 identity_rules 和默认参数，但不复制历史文章、画像、反馈或生成记录。")
    add_callout(doc, "参数优先级", "生成时按“系统默认值 → 角色默认值 → 本次请求显式值”合并。本次请求优先，响应中的 effective_request 是最终实际使用的完整参数。")

    add_heading(doc, "5. 投喂过去的文章", 1)
    add_body(doc, "投喂的目标不是让模型死记原文，而是建立可检索的真实写作样本。每篇文章应作为独立 Post 入库。不要只保存标题或链接。")
    add_heading(doc, "5.1 方式一：逐篇添加", 2)
    add_body(doc, "适合文章不多、希望精确标注每篇元数据的情况。使用 POST /posts：")
    add_code(doc, '{\n  "role_id": 3,\n  "title": "为什么品牌不能购买 AI 第一推荐",\n  "text": "这里填写完整历史文章正文……",\n  "platform": "LinkedIn",\n  "language": "zh-CN",\n  "content_type": "风险教育",\n  "topic": "GEO, AI Visibility, 品牌传播",\n  "tone": "克制, 教育型, 基于证据",\n  "authenticity": 5,\n  "published_at": "2026-06-15T12:00:00+00:00",\n  "source_name": "LinkedIn 历史文章"\n}')
    add_heading(doc, "5.2 元数据字段说明", 2)
    add_table(doc, ["字段", "用途", "示例"], [
        ("platform", "匹配发布环境", "LinkedIn、微信公众号、小红书、X"),
        ("language", "匹配语言并辅助双语画像", "zh-CN、en"),
        ("content_type", "匹配写作任务类型", "教育、打假、产品宣传、案例、活动"),
        ("topic", "描述文章讨论内容，可用逗号分隔", "GEO, AI Visibility, 品牌风险"),
        ("tone", "描述语气和判断方式", "克制, 教育型, 基于证据"),
        ("authenticity", "表示文章有多能代表角色本人", "1–5"),
        ("published_at", "计算较小的新鲜度分量", "ISO 8601 时间，可留空"),
        ("source_name", "追踪文章来源", "微信公众号导出、手工录入"),
    ], [1600, 4300, 3460], first_col_bold=True)
    add_heading(doc, "5.3 authenticity 如何评分", 2)
    add_table(doc, ["评分", "含义", "是否用于画像"], [
        ("5", "非常像本人，是最有代表性的作品", "是"),
        ("4", "整体像本人，可作为稳定画像样本", "是"),
        ("3", "普通内容，可用于检索但不应主导画像", "否"),
        ("2", "受客户、编辑或活动要求影响较大", "否"),
        ("1", "不代表本人，仅保留作业务参考", "否"),
    ], [1000, 6060, 2300], first_col_bold=True)
    add_callout(doc, "注意", "真实性不是流量分。高流量文章不一定最像本人；表现普通但充分体现判断方式的文章，往往更适合画像。")

    add_heading(doc, "5.4 方式二：批量上传", 2)
    add_body(doc, "推荐使用 POST /posts/bulk-upload，一次选择最多 30 个 PDF、TXT 或 Markdown 文件；单文件最大 10 MB、整批最大 50 MB。只需填写一次 role_id 和真实性，其余元数据可以留空。POST /posts/upload 与它共用同一套自动标注和去重逻辑。")
    add_body(doc, "推荐使用 Markdown，并通过 --- 分隔文章：")
    add_code(doc, "# 为什么品牌不能保证第一推荐\n\n第一篇文章完整正文……\n\n---\n\n# 隐藏文字会带来什么风险\n\n第二篇文章完整正文……\n\n---\n\n# AI Visibility 应该如何衡量\n\n第三篇文章完整正文……")
    add_body(doc, "系统会根据文件名和正文推断标题、平台、语言、内容类型、主题、语气与发布日期，并用规范化正文指纹跳过重复文章。批量数据库写入在一个事务中完成，写入异常时整批回滚。")
    add_heading(doc, "5.5 导入后纠错", 2)
    add_body(doc, "自动标注不确定时，使用 PATCH /posts/{post_id} 修正单篇；同一批文章标错时，使用 PATCH /posts/bulk 提交 role_id、post_ids 和 changes。批量修正会先验证全部文章属于该角色，任意一条不符合时整批不修改。")
    add_code(doc, '{\n  "role_id": 3,\n  "post_ids": [12, 13, 14],\n  "changes": {"platform": "小红书", "content_type": "项目复盘", "authenticity": 5}\n}')
    add_heading(doc, "5.6 语料质量检查", 2)
    for item in (
        "正文完整：不要只传标题、摘要或链接。",
        "角色正确：确认 role_id 没有填成其他账号。",
        "去除重复：同一篇文章的几十个小改版会挤占检索 Top-K。",
        "事实时效：历史数字可能过期，不能自动当成新文章事实。",
        "代表性分布：同时保留教育、宣传、案例、观点等不同类型的真实作品。",
        "扫描 PDF：当前只读取 PDF 文本层；纯图片 PDF 需要先做 OCR。",
    ):
        add_list_item(doc, item)

    doc.add_page_break()
    add_heading(doc, "6. 构建作者画像", 1)
    add_body(doc, "导入高真实性文章后，调用 POST /profile/rebuild?role_id=3。系统只选择 retrieval_status=ACTIVE 且 authenticity 为 4–5 的文章，最多使用 30 篇，并生成中英文画像。")
    add_body(doc, "画像包含价值观、voice、openings、rhythm、CTA 和 avoid。每次重建会保存新版本，不覆盖旧版本。使用 GET /profile?role_id=3 查看最新版本。")
    add_callout(doc, "常见报错", "如果返回“At least one post with authenticity >= 4 is required”，说明该角色没有高真实性样本。请先导入至少一篇 authenticity=4 或 5 的完整文章。")
    add_heading(doc, "画像与 RAG 的区别", 2)
    add_table(doc, ["组件", "稳定性", "示例"], [
        ("作者画像", "相对稳定", "喜欢克制判断、短段落、低压力 CTA"),
        ("RAG", "每次任务变化", "本次讨论 AI 第一推荐时，取回相关旧文"),
        ("写作任务", "每次任务变化", "这次写 LinkedIn 中文教育短文，给品牌负责人"),
    ], [1800, 1800, 5760], first_col_bold=True)

    add_heading(doc, "7. RAG 检索怎么工作", 1)
    add_body(doc, "检索器先把 topic、format、tone、goal、audience、platform 和 language 拼成查询，再与当前角色的文章比较。proof_points 不参与风格检索，因为事实白名单和风格参考属于不同信任层级。")
    add_heading(doc, "7.1 默认混合评分", 2)
    add_table(doc, ["分量", "权重", "解决的问题"], [
        ("字符 TF-IDF", "60%", "正文和主题是否具有相似字符片段"),
        ("关键词覆盖", "18%", "查询核心词是否出现在文章中"),
        ("元数据匹配", "14%", "平台、语言、内容类型和语气是否一致"),
        ("真实性", "6%", "是否优先使用更能代表本人的文章"),
        ("新鲜度", "2%", "在相关性接近时轻微偏向较新内容"),
    ], [2600, 1200, 5560], first_col_bold=True)
    add_body(doc, "初排后再使用 MMR 去重，防止高度相似的文章占满所有参考位置。")
    add_heading(doc, "7.2 生成前先检查 /retrieve", 2)
    add_code(doc, '{\n  "role_id": 3,\n  "topic": "为什么不能保证品牌成为 ChatGPT 第一推荐",\n  "platform": "LinkedIn",\n  "language": "zh-CN",\n  "format": "风险教育",\n  "tone": "克制, 基于证据",\n  "goal": "澄清市场误区",\n  "audience": "品牌市场负责人",\n  "top_k": 4\n}')
    add_body(doc, "重点查看 final_score、各分量得分、reason 和正文是否真的适合参考。如果结果不理想，依次检查 topic、content_type、tone、language、platform 和任务描述。")
    add_callout(doc, "不要这样做", "不要为了凑满 Top-K 而接受明显无关文章，也不要一开始就盲目修改检索权重。多数早期问题来自标签含糊、语料重复或任务描述太宽。")
    add_heading(doc, "7.3 检索纠错与文章停用", 2)
    add_body(doc, "某篇命中文章只是不适合本次任务时，向 POST /generations/{run_id}/retrieval-feedback 提交 NOT_RELEVANT。当前该动作返回 effect=RECORDED_ONLY，只留下任务级审计，不会暗中修改排序或真实性。")
    add_body(doc, "文章已经过时或不再代表本人时，向 POST /posts/{post_id}/retrieval-status 提交 RETIRE；人工确认恢复时提交 RESTORE。只有 ACTIVE 文章能进入后续 RAG 和画像重建，所有停用与恢复动作可通过 /posts/{post_id}/retrieval-actions 查询。")

    add_heading(doc, "8. 提交生成任务", 1)
    add_body(doc, "调用 POST /generate。角色已配置 default_generate_params 时，通常只需要提交 role_id、topic 和 proof_points；需要变化的字段在本次请求中覆盖。建议第一次生成 3 个候选版本。")
    add_code(doc, '{\n  "role_id": 3,\n  "topic": "品牌能否保证在 ChatGPT 中成为第一推荐？",\n  "platform": "LinkedIn",\n  "language": "zh-CN",\n  "format": "教育型短文",\n  "goal": "澄清常见误区并建立可信度",\n  "audience": "品牌市场负责人",\n  "tone": "克制、教育型、基于证据",\n  "length": "300-500字",\n  "banned_phrases": ["绝对保证", "颠覆行业", "不可逆趋势"],\n  "proof_points": [\n    "模型回答会受到问题表达、上下文和可用信息影响",\n    "不同问题可能得到不同推荐结果"\n  ],\n  "cta": "邀请读者分享他们观察 AI 品牌可见性的方式。",\n  "candidates": 3\n}')
    add_heading(doc, "8.1 生成字段说明", 2)
    add_table(doc, ["字段", "填写原则"], [
        ("topic", "写清要回答的具体问题，不要只写“AI”或“品牌”"),
        ("platform", "决定发布环境和内容习惯"),
        ("language", "使用 zh-CN 或 en 等稳定值"),
        ("format", "教育短文、产品发布、案例、打假、线程等"),
        ("goal", "读者看完后应该理解、相信或采取什么行动"),
        ("audience", "明确职位、知识水平和主要关切"),
        ("tone", "本次语气；不应与角色画像冲突"),
        ("length", "使用可理解的范围，例如 300–500 字"),
        ("banned_phrases", "本次明确禁止出现的词句"),
        ("proof_points", "本次允许公开使用的事实白名单"),
        ("cta", "希望使用的低压力或明确行动邀请"),
        ("candidates", "1–3 个候选版本"),
    ], [2200, 7160], first_col_bold=True)
    add_callout(doc, "核对实际参数", "查看响应中的 effective_request。它展示系统默认值、角色默认值和本次显式参数合并后的最终结果；HTTP、CLI 和代码直接调用现在使用同一套解析逻辑。")
    add_heading(doc, "8.2 proof_points 怎么写", 2)
    add_body(doc, "proof_points 应该是可确认、可公开、与本次任务有关的事实。")
    add_table(doc, ["好的 proof point", "不合格的写法"], [
        ("产品目前支持中英文内容分析", "产品能力很强"),
        ("本次测试使用了 20 篇公开文章", "测试规模很大"),
        ("模型回答会受到问题表达和上下文影响", "排名一定会变化"),
        ("该案例已经获得客户书面公开许可", "客户都很满意"),
    ], [4680, 4680])
    add_callout(doc, "数字门禁", "生成结果中的具体数字如果没有出现在 proof_points，本地规则会把它标成问题。无数字的因果结论仍需人工和评审器检查。")

    add_heading(doc, "9. 生成、评审与改写", 1)
    add_body(doc, "生成器只负责写初稿。初稿先经过本地规则，再由独立模型评审；评审不通过时只针对问题修改，最多两轮。")
    add_table(doc, ["检查层", "主要检查"], [
        ("本地硬规则", "未经允许的数字、禁用表达、与历史样本过度相似、内容过短"),
        ("独立模型评审", "角色一致性、任务完成度、AI 模板感、照抄、虚构、营销强度、CTA"),
        ("定向改写", "只修复 issues 中的问题，不引入新事实"),
        ("停止条件", "PASS，或完成最多两次改写"),
    ], [2200, 7160], first_col_bold=True)
    add_body(doc, "每次生成会返回 run_id、初稿、终稿、状态、评审记录、检索文章、effective_request 和按顺序保存的运行 steps。使用 GET /generations/{run_id} 可以回看完整运行资料。")
    add_heading(doc, "9.1 人工审批闭环", 2)
    add_body(doc, "成功生成后不会自动发布，而是进入 PENDING_REVIEW。使用 GET /review-inbox?role_id=3&status=PENDING_REVIEW&limit=100&offset=0 分页查看；列表中的 review_summary 直接说明审核次数、改写轮数、累计问题和阻塞原因。")
    add_body(doc, "向 POST /generations/{run_id}/review-actions 提交 EDIT、APPROVE 或 REJECT。EDIT 保存人工终稿并继续等待批准；APPROVE 确认发布前版本；REJECT 必须填写原因。每个动作都保留前后文本和 diff，编辑已批准内容会重新进入待审。")

    add_heading(doc, "10. 保存人工反馈", 1)
    add_body(doc, "真正让系统越来越像本人的数据，是“Agent 初稿—本人终稿—修改原因”。调用 POST /feedback：")
    add_code(doc, '{\n  "role_id": 3,\n  "task_summary": "LinkedIn 中文文章：解释为什么不能保证 AI 第一推荐",\n  "draft": "这里粘贴 Agent 原始版本",\n  "final_text": "这里粘贴本人修改后的最终版本",\n  "reason": "删除绝对化判断；缩短开头；把强销售 CTA 改成低压力讨论邀请；减少机械排比。",\n  "similarity_rating": 4\n}')
    add_heading(doc, "10.1 修改原因应该怎么写", 2)
    for item in (
        "指出具体问题：例如“开头铺垫太长”，不要只写“不像我”。",
        "说明修改方向：例如“把绝对结论改成带条件的判断”。",
        "记录事实原因：例如“该客户案例没有公开授权，全部删除”。",
        "记录结构偏好：例如“减少三段式排比，改成长短句交替”。",
        "记录 CTA 偏好：例如“不要预约演示，改成邀请分享观察”。",
    ):
        add_list_item(doc, item)
    add_body(doc, "新反馈先进入 CANDIDATE，不会立即影响生成。人工确认它代表稳定偏好后，向 POST /feedback/{feedback_id}/admission 提交 ADMIT；一次性活动要求或错误修改提交 REJECT。只有 ADMITTED 反馈会进入后续检索。")
    add_heading(doc, "10.2 批量准入反馈", 2)
    add_body(doc, "候选反馈较多时，使用 POST /feedback/admission/bulk，一次最多处理 100 条。所有 ID 必须属于指定角色且仍为 CANDIDATE，任一条不符合时整批回滚。")
    add_code(doc, '{\n  "role_id": 3,\n  "feedback_ids": [21, 22, 23],\n  "action": "ADMIT",\n  "reason": "均代表稳定的写作偏好"\n}')

    doc.add_page_break()
    add_heading(doc, "11. 一次完整示例", 1)
    add_heading(doc, "场景", 2)
    add_body(doc, "你要为个人 LinkedIn 写一篇中文文章，解释“品牌是否能购买 ChatGPT 第一推荐”。")
    steps = [
        ("确认角色", "GET /roles，确认个人账号的 role_id=3。"),
        ("准备语料", "用 /posts/bulk-upload 导入真实文章，检查自动标签，代表性文章设为 4–5。"),
        ("构建画像", "POST /profile/rebuild?role_id=3，检查中文画像是否符合本人。"),
        ("检查检索", "POST /retrieve，确认前几篇确实讨论第一推荐、承诺风险或 AI 可见性。"),
        ("准备事实", "把“回答受问题表达、上下文和可用信息影响”等可公开事实写入 proof_points。"),
        ("生成候选", "POST /generate，candidates=3。"),
        ("人工审批", "在 review-inbox 中编辑候选，确认事实后 APPROVE 或说明原因 REJECT。"),
        ("保存反馈", "POST /feedback 写明删改原因，再决定 ADMIT 或 REJECT。"),
        ("回看记录", "GET /generations/{run_id}，确认检索、评审和终稿均已保存。"),
    ]
    add_table(doc, ["步骤", "操作"], steps, [1800, 7560], first_col_bold=True)
    add_callout(doc, "成功标准", "不是文章一次生成就不用改，而是检索参考正确、事实没有越界、修改时间比通用模型明显减少，并且反馈能让后续任务持续改善。")

    add_heading(doc, "12. 日常运营方法", 1)
    add_heading(doc, "12.1 每次写作", 2)
    for item in (
        "确认角色和本次公开事实。",
        "先检索，后生成。",
        "生成 1–3 个候选，不自动发布。",
        "人工核对事实、身份和敏感信息。",
        "保存有代表性的修改反馈，并只准入稳定偏好。",
    ):
        add_list_item(doc, item, ordered=True)
    add_heading(doc, "12.2 每周检查", 2)
    for item in (
        "哪些主题经常出现 NOT_RELEVANT 记录？",
        "哪些标签过宽、过细或写法不一致？",
        "是否出现大量重复文章？",
        "哪些反馈原因反复出现？",
        "是否存在应 RETIRE 的过期文章或不应公开语料？",
    ):
        add_list_item(doc, item)
    add_heading(doc, "12.3 什么时候重建画像", 2)
    add_body(doc, "当新增一批高真实性文章、角色定位发生调整、中文/英文风格有明显变化时重建。不要每写一篇文章就重建画像；单次修改更适合进入 feedback。")

    add_heading(doc, "13. 测试与排错", 1)
    add_heading(doc, "13.1 离线测试", 2)
    add_code(doc, "$env:LLM_MODE=\"mock\"\npython -m unittest discover -s tests -v\npython -m scripts.smoke_test")
    add_body(doc, "当前完整测试共 57 项，覆盖字符 n-gram、混合排序、角色隔离、批量导入与事务回滚、默认参数、反馈准入、人工审批、文章停用恢复、数字门禁、评测和完整运行轨迹。mock 模式不会把文章发给外部模型。")
    add_heading(doc, "13.2 常见问题", 2)
    add_table(doc, ["现象", "优先检查"], [
        ("/health 无法访问", "服务是否启动、端口是否为 8000、终端是否有报错"),
        ("模型鉴权失败", ".env 是否存在、供应商是否正确、密钥是否过期或无余额"),
        ("画像无法重建", "是否存在 authenticity>=4 的完整文章"),
        ("检索文章不相关", "role_id、topic、content_type、tone、language、重复语料"),
        ("输出出现错误数字", "是否把数字放进 proof_points；人工核验评审结果"),
        ("输出太像旧文", "检查重复语料、Top-K、MMR 结果与相似度问题"),
        ("PDF 没有正文", "是否为扫描 PDF；先 OCR 再上传"),
        ("不同账号串味", "检查创建和上传时使用的 role_id；正式多用户版还需租户权限"),
    ], [3000, 6360], first_col_bold=True)

    add_heading(doc, "14. 数据与安全边界", 1)
    for item in (
        "当前 SQLite 文件是 data/voice.db，可理解为可查询的本地数据表。",
        "MVP 通过 role_id 隔离数据，但还没有登录、权限和 tenant_id；正式多用户上线前必须补齐。",
        "生成内容必须人工确认后再发布，特别是客户案例、法律、医疗、金融和监管话题。",
        "不要把 API Key、个人隐私、未公开客户资料或合同内容放进历史文章。",
        "历史文章中的指令文本只是数据，不能覆盖角色边界和系统规则。",
        "定期备份数据库，并建立删除、导出和保留期限。",
    ):
        add_list_item(doc, item)

    add_heading(doc, "15. API 快速索引", 1)
    add_table(doc, ["方法", "路径", "用途"], [
        ("GET", "/health", "检查服务、供应商、模型和检索模式"),
        ("GET", "/roles", "列出角色"),
        ("POST/PATCH", "/roles · /roles/{id}", "创建或更新角色与默认任务参数"),
        ("POST", "/roles/{id}/clone", "复制角色规则和默认参数，不复制记忆"),
        ("GET", "/posts?role_id=…", "查看角色文章"),
        ("POST", "/posts", "逐篇添加文章"),
        ("POST", "/posts/bulk-upload", "批量上传、自动标注、去重并事务写入"),
        ("PATCH", "/posts/bulk", "批量纠正文章元数据"),
        ("POST", "/posts/{id}/retrieval-status", "停用或恢复文章的检索资格"),
        ("POST", "/profile/rebuild?role_id=…", "重建作者画像"),
        ("GET", "/profile?role_id=…", "查看最新画像"),
        ("POST", "/retrieve", "单独检查 RAG 结果"),
        ("POST", "/generate", "执行生成、评审和改写"),
        ("GET/POST", "/feedback", "查询反馈或创建候选反馈"),
        ("POST", "/feedback/admission/bulk", "批量准入或拒绝候选反馈"),
        ("GET", "/review-inbox", "分页查看待人工审核内容与摘要"),
        ("POST", "/generations/{id}/review-actions", "编辑、批准或拒绝内容"),
        ("GET", "/generations/{run_id}", "回看一次生成运行"),
        ("POST", "/eval-runs", "对已有生成运行执行确定性评测"),
    ], [1100, 3400, 4860])

    add_heading(doc, "附录 A：当前项目文件", 1)
    add_table(doc, ["文件", "职责"], [
        ("app/main.py", "FastAPI 接口"),
        ("app/db.py", "SQLite 数据库与角色隔离"),
        ("app/documents.py", "PDF/TXT/Markdown 提取与文章拆分"),
        ("app/ingestion.py", "自动元数据、正文指纹、统一上传流程"),
        ("app/retrieval.py", "混合检索、评分解释与 MMR"),
        ("app/agent.py", "画像、生成、评审、改写与反馈编排"),
        ("app/llm.py", "DeepSeek Chat Completions 与 OpenAI Responses API"),
        ("app/models.py", "接口字段与校验"),
        ("app/evals/", "确定性评测、运行比较和结果报告"),
        ("scripts/smoke_test.py", "不修改正式数据库的端到端冒烟测试"),
        ("RAG_DESIGN.md", "权重、评测、升级与生产化设计"),
    ], [3000, 6360], first_col_bold=True)

    add_heading(doc, "附录 B：上线前检查清单", 1)
    for item in (
        "角色名称、描述和身份边界已经确认。",
        "高真实性样本足够且没有混入其他账号。",
        "检索结果在代表任务上经过人工检查。",
        "proof_points 的来源、公开权限和时效已经确认。",
        "生成结果不会自动发布。",
        "API Key 没有进入 Git、日志或截图。",
        "正式多用户版本已经加入认证、租户隔离和审计。",
        "模型、提示词、画像和检索器版本可以回溯。",
        "数据库具备备份、恢复和删除流程。",
    ):
        add_list_item(doc, item)

    # Core properties and save.
    doc.core_properties.title = "个性化社交媒体写作 Agent 详细使用手册"
    doc.core_properties.subject = "角色创建、文章投喂、RAG 检索、生成评审与反馈闭环"
    doc.core_properties.author = "SocialPostEditor"
    doc.core_properties.keywords = "Agent, RAG, 社交媒体, 写作, 操作手册"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
