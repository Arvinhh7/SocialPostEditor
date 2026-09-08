from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "BEGINNER_USAGE_GUIDE.md"
OUTPUT = ROOT / "个性化社交媒体写作Agent_零基础使用手册.docx"
QA_DIR = ROOT / ".qa" / "beginner_manual"
REFERENCE = Path(
    r"C:\Users\hao32\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-design-report\assets\reference.docx"
)

INK = "101820"
BLUE = "245B78"
PALE_BLUE = "EAF2F6"
PALE_GRAY = "F4F5F6"
GRID = "D6DBDE"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), GRID)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, size: float | None = None, bold: bool | None = None, code: bool = False) -> None:
    name = "Consolas" if code else "Helvetica Neue"
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(INK)


def set_paragraph_basics(paragraph, after=6, before=0, line=1.15) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.line_spacing = line
    p_pr = paragraph._p.get_or_add_pPr()
    widow = p_pr.find(qn("w:widowControl"))
    if widow is None:
        widow = OxmlElement("w:widowControl")
        p_pr.append(widow)


def add_inline(paragraph, text: str, size=10.5, bold=False) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size - 0.5, code=True)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold)


def add_body_paragraph(doc, text: str, *, bullet=False, number=False) -> None:
    paragraph = doc.add_paragraph()
    if bullet or number:
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), "1" if bullet else "2")
        num_pr.extend((ilvl, num_id))
        paragraph._p.get_or_add_pPr().append(num_pr)
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    add_inline(paragraph, text)
    set_paragraph_basics(paragraph, after=4 if (bullet or number) else 7)


def add_heading(doc, text: str, level: int) -> None:
    style = "Heading 1" if level == 2 else "Heading 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = level == 2 and text in {
        "第一部分 准备 Windows 环境",
        "第四部分 完整使用流程",
        "常见问题排查",
    }
    set_paragraph_basics(paragraph, after=7, before=16 if level == 2 else 11, line=1.0)
    run = paragraph.add_run(text)
    set_run_font(run, size=23 if level == 2 else 14, bold=True)


def add_code_block(doc, lines: list[str]) -> None:
    for index, line in enumerate(lines or [""]):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(5 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(5 if index == len(lines) - 1 else 0)
        paragraph.paragraph_format.keep_with_next = index < len(lines) - 1
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), PALE_GRAY)
        paragraph._p.get_or_add_pPr().append(shd)
        run = paragraph.add_run(line or " ")
        set_run_font(run, size=8.5, code=True)


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    usable = 6.45
    width = usable / columns
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, BLUE if row_index == 0 else (PALE_BLUE if row_index % 2 == 0 else WHITE))
            cell.text = ""
            paragraph = cell.paragraphs[0]
            set_paragraph_basics(paragraph, after=0, line=1.05)
            value = values[column_index] if column_index < len(values) else ""
            add_inline(paragraph, value, size=8.8, bold=row_index == 0)
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_markdown(doc, markdown: str) -> str:
    lines = markdown.splitlines()
    title = lines[0].removeprefix("# ").strip()
    index = 1
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
        elif re.match(r"^\d+\.\s+", line):
            add_body_paragraph(doc, re.sub(r"^\d+\.\s+", "", line), number=True)
        elif line.startswith("- "):
            add_body_paragraph(doc, line[2:], bullet=True)
        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = []
            for item in table_lines:
                cells = [cell.strip() for cell in item.strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                parsed.append(cells)
            add_table(doc, parsed)
            continue
        else:
            add_body_paragraph(doc, line)
        index += 1
    return title


def create_cover_image(path: Path) -> None:
    width, height = 1600, 1480
    image = Image.new("RGB", (width, height), "#F3F0E9")
    draw = ImageDraw.Draw(image)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    regular = ImageFont.truetype(str(font_path), 34) if font_path.exists() else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 26) if font_path.exists() else ImageFont.load_default()
    bold = ImageFont.truetype(str(bold_path), 46) if bold_path.exists() else regular
    draw.text((110, 95), "从历史文章到可审批终稿", fill="#101820", font=bold)
    draw.text((112, 165), "零基础完整工作路径", fill="#52616B", font=regular)
    labels = [
        ("01", "创建角色", "保存身份规则和常用参数"),
        ("02", "批量投喂", "上传 PDF TXT Markdown"),
        ("03", "建立画像", "提取真实写作偏好"),
        ("04", "RAG 检索", "找到相关历史表达"),
        ("05", "生成审核", "写作 Agent 自动检查改写"),
        ("06", "人工闭环", "审批并准入长期反馈"),
    ]
    card_w, card_h = 620, 270
    positions = [(100, 280), (880, 280), (100, 610), (880, 610), (100, 940), (880, 940)]
    for (number, label, note), (x, y) in zip(labels, positions):
        draw.rounded_rectangle((x, y, x + card_w, y + card_h), radius=26, fill="#FFFFFF", outline="#D0D8DC", width=3)
        draw.rounded_rectangle((x + 28, y + 30, x + 125, y + 127), radius=20, fill="#245B78")
        draw.text((x + 48, y + 54), number, fill="#FFFFFF", font=small)
        draw.text((x + 155, y + 38), label, fill="#101820", font=bold)
        draw.text((x + 155, y + 122), note, fill="#52616B", font=small)
        draw.line((x + 155, y + 190, x + 560, y + 190), fill="#DCE4E8", width=3)
    draw.text((110, 1340), "安全起点  Mock 模式验证后再连接真实模型", fill="#245B78", font=regular)
    image.save(path, quality=95)


def replace_docx_media(docx_path: Path, media_path: Path) -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        extracted = Path(temp_dir) / "package"
        with zipfile.ZipFile(docx_path, "r") as archive:
            archive.extractall(extracted)
        target = extracted / "word" / "media" / "image1.png"
        shutil.copyfile(media_path, target)
        rebuilt = Path(temp_dir) / "rebuilt.docx"
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as archive:
            for item in extracted.rglob("*"):
                if item.is_file():
                    archive.write(item, item.relative_to(extracted).as_posix())
        shutil.copyfile(rebuilt, docx_path)


def configure_styles(doc) -> None:
    normal = doc.styles["normal"]
    normal.font.name = "Helvetica Neue"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for style_name, size in (("Heading 1", 23), ("Heading 2", 14), ("Title", 38)):
        style = doc.styles[style_name]
        style.font.name = "Helvetica Neue"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)


def update_cover(doc, title: str) -> None:
    cover_title = doc.paragraphs[2]
    cover_title.text = title
    cover_title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    for run in cover_title.runs:
        set_run_font(run, size=38, bold=True)
    table = doc.tables[0]
    values = [
        "从环境安装到角色创建、文章投喂、RAG 检索、生成审核与反馈闭环",
        "",
        "SocialPostEditor\n当前 GitHub main 分支",
    ]
    for cell, value in zip(table.rows[0].cells, values):
        cell.text = value
        for paragraph in cell.paragraphs:
            set_paragraph_basics(paragraph, after=0, line=1.05)
            for run in paragraph.runs:
                set_run_font(run, size=10.5, bold=False)


def remove_template_body(doc) -> None:
    body = doc._element.body
    children = list(body)
    for child in children[5:-1]:
        body.remove(child)


def enable_field_updates(doc) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    QA_DIR.mkdir(parents=True, exist_ok=True)
    markdown = SOURCE.read_text(encoding="utf-8")
    document = Document(REFERENCE)
    configure_styles(document)
    remove_template_body(document)
    title = markdown.splitlines()[0].removeprefix("# ").strip()
    update_cover(document, title)
    parse_markdown(document, markdown)
    enable_field_updates(document)
    document.core_properties.title = title
    document.core_properties.subject = "SocialPostEditor 零基础安装与使用说明"
    document.core_properties.author = "SocialPostEditor"
    document.save(OUTPUT)
    cover_image = QA_DIR / "workflow-cover.png"
    create_cover_image(cover_image)
    replace_docx_media(OUTPUT, cover_image)
    print("Beginner DOCX created successfully")


if __name__ == "__main__":
    main()
