from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "个性化社交媒体写作Agent_详细使用手册.docx"


def main() -> None:
    doc = Document(PATH)
    errors: list[str] = []
    section = doc.sections[0]
    expected_inches = {
        "page_width": 8.5,
        "page_height": 11.0,
        "top_margin": 1.0,
        "bottom_margin": 1.0,
        "left_margin": 1.0,
        "right_margin": 1.0,
    }
    for name, expected in expected_inches.items():
        actual = getattr(section, name).inches
        if abs(actual - expected) > 0.01:
            errors.append(f"{name}: {actual} != {expected}")

    for index, table in enumerate(doc.tables, start=1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_w is None or tbl_w.get(qn("w:w")) != "9360":
            errors.append(f"table {index}: tblW")
        if tbl_ind is None or tbl_ind.get(qn("w:w")) != "120":
            errors.append(f"table {index}: tblInd")
        widths = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        if sum(widths) != 9360:
            errors.append(f"table {index}: grid sum {sum(widths)}")
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells):
                tc_w = cell._tc.tcPr.find(qn("w:tcW"))
                if tc_w is None or int(tc_w.get(qn("w:w"))) != widths[cell_index]:
                    errors.append(f"table {index} row {row_index} cell {cell_index + 1}: tcW")

    text = "\n".join(p.text for p in doc.paragraphs)
    for marker in ("TODO", "FIXME", "Lorem ipsum", "sk-"):
        if marker in text:
            errors.append(f"placeholder or secret marker: {marker}")
    heading_counts = {
        level: sum(1 for p in doc.paragraphs if p.style.name == f"Heading {level}")
        for level in (1, 2, 3)
    }
    numbered = sum(1 for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None)
    if numbered == 0:
        errors.append("no real numbered/list paragraphs")
    if errors:
        raise SystemExit("AUDIT FAILED\n" + "\n".join(errors))
    print(
        f"AUDIT PASS | paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} "
        f"headings={heading_counts} list_items={numbered} characters={len(text)} bytes={PATH.stat().st_size}"
    )


if __name__ == "__main__":
    main()
