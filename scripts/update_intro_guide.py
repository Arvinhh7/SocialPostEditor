from __future__ import annotations

from pathlib import Path

from docx import Document

from scripts.build_usage_guide import set_run_font, style_table


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "个性化社交媒体写作Agent_RAG项目入门指南.docx"
OUTPUT = SOURCE


REPLACEMENTS = {
    "完整链路是：": "完整链路是：创建或克隆角色 → 批量导入文章 → 自动拆分、标注和去重 → 建立画像 → 检索 ACTIVE 文章 → 生成与独立评审 → 人工编辑和审批 → 保存反馈 → 人工决定是否准入长期记忆。",
    "角色不是一个昵称": "角色不是一个昵称，而是一套独立的写作身份。每个角色都有自己的语料、作者画像、反馈、生成记录和 default_generate_params。需要同一作者管理另一个平台时可以克隆角色规则，但不会复制历史文章和长期记忆。",
    "上传 PDF 后": "上传 PDF、TXT 或 Markdown 后，系统会提取文字、拆分独立 Post，并根据文件名与正文推断标题、平台、语言、内容类型、主题、语气和发布日期。单文件与批量上传共用同一套去重和自动标注逻辑。",
    "目前角色 2 已入库": "语料数量会随实际导入变化，不应把固定篇数写进系统假设。完整正文才会入库；标题、空片段和不足 80 字的内容会被跳过或列入失败清单。",
    "每篇文章不是只使用正文": "每篇文章会把 content_type、topic、tone、platform、language、title 和 text 拼成检索文档。本次查询使用 topic、format、tone、goal、audience、platform 和 language；proof_points 属于事实白名单，不参与风格检索。",
    "系统先计算查询向量": "系统先计算字符 TF-IDF 相似度，再结合关键词覆盖、平台/语言/内容类型/语气匹配、真实性和新鲜度形成可解释分数，最后用 MMR 避免近重复文章占满 Top-K。",
    "当前排序公式": "当前 hybrid 公式  最终分数 = 字符 TF-IDF 60% + 关键词 18% + 元数据 14% + 真实性 6% + 新鲜度 2%。tfidf 兼容模式仍使用相似度加很小的真实性分量。",
    "调用 POST /profile/rebuild": "调用 POST /profile/rebuild?role_id=2 后，模型只分析 retrieval_status=ACTIVE 且 authenticity 为 4–5 的样本，最多使用 30 篇，分别总结中英文的价值观、voice、openings、rhythm、CTA 和 avoid。",
    "第一次输出只是初稿": "第一次输出只是初稿。本地规则和独立评审器会检查事实、角色一致性、模板感、照抄与营销强度；审核故障会 BLOCKED，不会降级成 PASS。成功生成后仍进入 PENDING_REVIEW，必须由人工编辑、批准或拒绝。",
    "最有价值的数据往往不是": "最有价值的数据往往是“AI 初稿 → 本人终稿 → 修改原因”。新反馈先进入 CANDIDATE，只有人工 ADMIT 后才会被后续任务检索；一次性要求或错误反馈应 REJECT，避免污染长期记忆。",
    "当前使用 SQLite": "当前使用 SQLite，本地文件为 data/voice.db。数据库除角色、文章、画像和生成记录外，还保存反馈准入、生成步骤、人工审批、检索负反馈、文章停用恢复和评测审计。",
    "当前只有少量文章": "当前语料规模适合使用本地、可解释的混合检索。直接引入 Qdrant、Milvus 或 pgvector 会增加部署成本；应先用固定评测案例确认现有检索的真实瓶颈。",
    "专业版可以同时使用": "当前版本已经同时使用字符 TF-IDF、关键词、元数据、真实性、新鲜度和 MMR。下一步只有在固定评测证明语义召回不足时，才考虑增加 Embedding 或向量数据库。",
    "示例权重": "当前实际权重  字符 TF-IDF 60% + 关键词覆盖 18% + 元数据匹配 14% + 真实性 6% + 新鲜度 2%；权重变化应通过 eval-cases 和 eval-runs 比较验证。",
    "1. 配置模型 API Key": "1. 先保持 LLM_MODE=mock，启动 FastAPI 并用 57 项测试和 smoke test 验证本地链路；需要真实生成时再配置 DeepSeek 或 OpenAI。",
    "2. 查看 GET /roles": "2. 查看 GET /roles；创建新角色时同时设置 identity_rules 和长期复用的 default_generate_params。",
    "3. 调用 POST /profile/rebuild": "3. 用 POST /posts/bulk-upload 导入真实文章，检查自动标签，必要时通过单条或批量 PATCH 修正，再重建画像。",
    "4. 选择 10 个真实选题": "4. 先调用 /retrieve 检查命中，再用 /generate 生成；通过 effective_request 核对最终生效参数。",
    "5. 由本人修改": "5. 在 review-inbox 中查看审核摘要，人工 EDIT、APPROVE 或 REJECT，系统不会自动发布。",
    "6. 观察哪些检索样本": "6. 保存初稿、终稿和修改原因，再人工 ADMIT 或 REJECT；错误命中记录 NOT_RELEVANT，过时文章使用 RETIRE。",
    "7. 积累数百篇语料后": "7. 建立 eval-cases，对已保存的 generation_run 执行确定性评测；只有指标证明需要时才升级 Embedding。",
}


def replace_table(table, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    style_table(table, first_col_bold=True)


def build() -> None:
    doc = Document(SOURCE)
    for paragraph in doc.paragraphs:
        for prefix, replacement in REPLACEMENTS.items():
            if paragraph.text.startswith(prefix):
                paragraph.text = replacement
                if paragraph.style.name == "Normal":
                    for run in paragraph.runs:
                        set_run_font(run, size=11)
                break

    replace_table(doc.tables[2], ["层次", "职责", "通俗类比"], [
        ("角色层", "身份边界与默认任务模板", "员工身份与品牌守则"),
        ("语料层", "真实文章、元数据、真实性和 ACTIVE/RETIRED 状态", "资料室"),
        ("画像层", "从高真实性 ACTIVE 样本总结稳定表达", "写作习惯卡"),
        ("检索层", "混合评分、解释和 MMR 去重", "按任务找资料"),
        ("生成层", "Evidence Contract、初稿、评审与定向改写", "受控写作流水线"),
        ("人工审批层", "编辑、批准、拒绝并保留 diff", "发布前总编"),
        ("反馈与评测层", "经验准入、检索纠错和确定性指标", "可审计的改进闭环"),
    ])
    replace_table(doc.tables[3], ["角色字段", "作用"], [
        ("name", "可识别的角色或账号名称"),
        ("description", "受众、主题和业务范围"),
        ("identity_rules", "不得虚构的身份、事实、承诺和合规边界"),
        ("default_generate_params", "长期复用的平台、语言、受众、语气、长度和禁用词"),
        ("role_id", "文章、画像、反馈和生成记录的隔离键"),
    ])
    replace_table(doc.tables[4], ["保存字段", "检索或生成中的用途"], [
        ("text / title", "正文与标题，是最主要的内容和风格参考"),
        ("platform / language", "匹配发布环境与语言"),
        ("content_type / topic / tone", "匹配任务类型、主题和表达方式"),
        ("authenticity", "4–5 可进入画像；1–5 均可辅助检索排序"),
        ("retrieval_status", "只有 ACTIVE 可进入 RAG 和画像；RETIRED 保留审计但不召回"),
        ("published_at / source_name", "提供轻量新鲜度和来源追踪"),
    ])
    replace_table(doc.tables[7], ["分量", "权重", "作用", "是否可解释"], [
        ("字符 TF-IDF", "60%", "中英文字符片段相似度", "是"),
        ("关键词覆盖", "18%", "核心词命中程度", "是"),
        ("元数据", "14%", "平台、语言、类型和语气匹配", "是"),
        ("真实性", "6%", "轻度偏向代表性文章", "是"),
        ("新鲜度", "2%", "相关性接近时轻度偏新", "是"),
    ])
    replace_table(doc.tables[8], ["因素", "当前状态", "说明"], [
        ("角色与状态过滤", "已使用", "只搜索同一 role_id 下的 ACTIVE 文章"),
        ("正文与主题", "已使用", "字符 TF-IDF 与关键词覆盖"),
        ("平台、语言、类型、语气", "已使用", "形成 14% 元数据分量"),
        ("真实性与新鲜度", "已使用", "只做小幅重排，不覆盖相关性"),
        ("MMR", "已使用", "减少近重复参考"),
        ("Embedding", "尚未使用", "需要由固定评测证明收益后再引入"),
    ])
    replace_table(doc.tables[12], ["数据表", "主要内容", "作用域"], [
        ("roles / posts / profiles", "角色配置、文章和画像版本", "role_id"),
        ("feedback / feedback_admission_actions", "候选反馈和准入审计", "role_id / feedback_id"),
        ("generation_runs / generation_run_steps", "生成结果和逐步轨迹", "generation_run_id"),
        ("review_actions", "人工编辑、批准、拒绝和 diff", "generation_run_id"),
        ("retrieval_feedback", "任务级 NOT_RELEVANT 记录", "generation_run_id + post_id"),
        ("post_retrieval_actions", "文章 RETIRE / RESTORE 审计", "post_id"),
        ("eval_cases / eval_runs", "固定案例、指标和比较结果", "role_id / eval_case_id"),
    ])
    replace_table(doc.tables[13], ["语料规模", "推荐方案"], [
        ("10–500 篇", "当前可解释混合检索 + 固定评测"),
        ("500–5,000 篇", "先评测，再考虑 Embedding + 本地向量索引"),
        ("更大或多租户", "向量数据库、tenant_id、权限与可观测性"),
    ])
    replace_table(doc.tables[15], ["文件", "职责"], [
        ("app/main.py", "FastAPI 路由与 HTTP 错误映射"),
        ("app/agent.py", "画像、参数解析、检索、生成、评审和改写"),
        ("app/retrieval.py", "混合评分、解释和 MMR"),
        ("app/ingestion.py", "自动元数据、正文指纹和统一上传流程"),
        ("app/db.py", "SQLite、事务、状态机与审计记录"),
        ("app/reviewer.py", "独立只读审核器"),
        ("app/evals/", "确定性指标、运行和比较"),
        ("tests/", "57 项离线回归测试"),
    ])

    if not any("当前闭环操作" in paragraph.text for paragraph in doc.paragraphs):
        anchor = next(p for p in doc.paragraphs if p.text.startswith("14. 数据库里保存什么"))
        heading = anchor.insert_paragraph_before("13.1 当前闭环操作", style="Heading 2")
        body = anchor.insert_paragraph_before(
            "生成响应会返回 effective_request 和 run_id。人工先在 review-inbox 查看 review_summary，再对生成记录执行 EDIT、APPROVE 或 REJECT；系统不会自动发布内容。"
        )
        body2 = anchor.insert_paragraph_before(
            "反馈创建后是 CANDIDATE，只有 ADMIT 才进入后续生成。NOT_RELEVANT 目前只做任务级记录；RETIRE / RESTORE 则直接控制文章是否能进入未来 RAG 和画像。"
        )
        heading2 = anchor.insert_paragraph_before("13.2 当前闭环操作的接口", style="Heading 2")
        body3 = anchor.insert_paragraph_before(
            "核心接口包括 /review-inbox、/generations/{run_id}/review-actions、/feedback/{feedback_id}/admission、/feedback/admission/bulk、/generations/{run_id}/retrieval-feedback 和 /posts/{post_id}/retrieval-status。"
        )
        for paragraph in (body, body2, body3):
            for run in paragraph.runs:
                set_run_font(run, size=11)
        _ = (heading, heading2)

    read_map = next(p for p in doc.paragraphs if p.text == "阅读地图")
    if not any("版本说明" in p.text for p in doc.paragraphs[:10]):
        version = read_map.insert_paragraph_before(
            "版本说明  本指南已按 2026-09-04 当前代码更新，覆盖 28 个 API 路径和 57 项自动化测试。"
        )
        for run in version.runs:
            set_run_font(run, size=10, color="596775", italic=True)

    doc.core_properties.subject = "当前代码对应的角色、RAG、生成、审批、反馈准入与评测入门指南"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
